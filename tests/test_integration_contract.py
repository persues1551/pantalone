from __future__ import annotations

import importlib.util
import json
import os
import re
import subprocess
import sys
import tomllib
from pathlib import Path

import pytest
import yaml
from docx import Document


ROOT = Path(__file__).resolve().parents[1]
CAPABILITY_MANIFEST = ROOT / "references/external-capabilities.yaml"
ACTIVE_DOCS = [
    ROOT / "README.md",
    ROOT / "SKILL.md",
    ROOT / "SOUL.md",
    ROOT / "router.md",
    ROOT / "workflow_v4_unified.md",
    *sorted((ROOT / "subagents").glob("*.md")),
    *sorted((ROOT / "rules").glob("*.md")),
    *sorted((ROOT / "templates").rglob("*.md")),
]
CORE_REFERENCES = [
    ROOT / f"references/{name}"
    for name in [
        "building-strategy-framework.md",
        "deep-stock-research-unified.md",
        "dragon-leader-database.md",
        "hk-stock-data-api.md",
        "hk-stock-trading-guide.md",
        "layer34-selection-timing-bridge.md",
        "mega-ipo-impact-analysis.md",
        "ocifq-framework.md",
        "patch-event-tracker.md",
        "sector-correlation-database.md",
        "sector-screening-workflow.md",
        "supply-chain-mapping-framework.md",
        "value-chain-analysis.md",
        "version-detection-framework.md",
    ]
]


def tracked_text_files():
    result = subprocess.run(
        ["git", "ls-files", "-z"],
        cwd=ROOT,
        check=True,
        capture_output=True,
    )
    for raw_path in result.stdout.split(b"\0"):
        if not raw_path:
            continue
        path = ROOT / os.fsdecode(raw_path)
        if not path.is_file():
            continue
        try:
            path.read_text(encoding="utf-8")
        except (UnicodeDecodeError, OSError):
            continue
        yield path


def load_schemas():
    path = ROOT / "subagents/schemas.py"
    spec = importlib.util.spec_from_file_location("pantalone_integration_schemas", path)
    assert spec is not None
    module = importlib.util.module_from_spec(spec)
    assert spec.loader is not None
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


def load_tencent_quote_parser():
    path = ROOT / "scripts/tencent_quote_parser.py"
    spec = importlib.util.spec_from_file_location("pantalone_tencent_quote_parser", path)
    assert spec is not None
    module = importlib.util.module_from_spec(spec)
    assert spec.loader is not None
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


def load_sim_integrator():
    path = ROOT / "scripts/amadeus_sim_integrate.py"
    spec = importlib.util.spec_from_file_location("pantalone_sim_integrator", path)
    assert spec is not None
    module = importlib.util.module_from_spec(spec)
    assert spec.loader is not None
    spec.loader.exec_module(module)
    return module


def test_active_execution_contract_has_no_private_machine_paths():
    patterns = [
        re.escape("~/" + ".hermes"),
        re.escape("/" + "home" + "/" + "ubuntu"),
        re.escape("/" + "root" + "/"),
        re.escape("/" + "Users" + "/"),
    ]
    findings = []
    for path in ACTIVE_DOCS:
        text = path.read_text()
        for pattern in patterns:
            if re.search(pattern, text):
                findings.append((str(path.relative_to(ROOT)), pattern))
    assert findings == []


def test_external_runtime_paths_are_declared_optional_and_fail_closed():
    text = "\n".join(path.read_text() for path in [*ACTIVE_DOCS, *CORE_REFERENCES])
    external = sorted(
        {
            match.rstrip("/.,;:)")
            for match in re.findall(r"\$HERMES_HOME/([A-Za-z0-9_./-]+)", text)
            if match.rstrip("/.,;:)") != "..."
        }
    )
    dependency_suffixes = (".py", ".md", "/SKILL.md")
    dependencies = [path for path in external if path.endswith(dependency_suffixes)]
    manifest = yaml.safe_load(CAPABILITY_MANIFEST.read_text())
    declared = manifest["capabilities"]
    undeclared = [path for path in dependencies if path not in declared]
    assert undeclared == []
    for path in dependencies:
        item = declared[path]
        assert item["required"] is False, path
        assert item["fallback"].strip(), path

    for name in [
        "market_data.md",
        "technical.md",
        "risk.md",
        "capital.md",
        "theme.md",
        "research.md",
        "financial.md",
        "macro.md",
    ]:
        content = (ROOT / "subagents" / name).read_text()
        assert "HERMES_HOME" in content
        assert "存在" in content
        assert any(word in content for word in ("缺失", "失败", "不可用"))


def test_unsafe_machine_bound_helpers_are_not_shipped():
    assert not (ROOT / "scripts/opencli_server.py").exists()
    assert not (ROOT / "scripts/verify_predictions.py").exists()


def test_workflow_registry_paths_and_claims_are_truthful():
    path = ROOT / "references/workflow-registry.yaml"
    registry = yaml.safe_load(path.read_text())
    assert isinstance(registry, dict)
    assert isinstance(registry.get("version"), str)
    assert "skills" in registry and "workflows" in registry

    missing = []
    for name, item in registry["skills"].items():
        target = item.get("path", "")
        if target.startswith("$HERMES_HOME/"):
            exists = item.get("required") is False and bool(item.get("fallback"))
        else:
            exists = (ROOT / target).exists()
        if not exists:
            missing.append((name, target))
    assert missing == []

    raw = path.read_text()
    assert "新增技能只需在此注册，自动接入" not in raw
    assert "自动生成/更新cron" not in raw
    assert "不是执行引擎" in raw
    assert "不会自动运行" in raw
    assert "仅修改本文件不会产生任何运行行为" in raw

    for workflow_name, workflow in registry["workflows"].items():
        available = set(workflow.get("inputs", []))
        for step_name in workflow.get("steps", []):
            assert step_name in registry["skills"], (workflow_name, step_name)
            step = registry["skills"][step_name]
            unresolved = set(step.get("consumes", [])) - available
            assert unresolved == set(), (workflow_name, step_name, sorted(unresolved))
            available.update(step.get("produces", []))


def test_core_schemas_instantiate_and_render():
    m = load_schemas()
    market = m.MarketDataReport(
        emotion={"score": 50, "level": "中性"},
        sector_flow={},
        market_filter={"level": "震荡", "position_pct": 40},
        northbound={"bias": "净流出"},
    )
    technical = m.TechnicalAnalysisReport(stocks={"600000": {"price": 10.0}})
    financial = m.FinancialReport(
        financial_score=70,
        revenue_trend="稳定",
        profit_trend="稳定",
        cashflow_quality="正常",
        peer_comparison="接近行业中位数",
    )
    theme = m.ThemeAnalysisReport(theme_score=60, theme_bias="中性")
    macro = m.MacroAnalysisReport(
        macro_analysis={
            "monetary_policy": "中性",
            "fiscal_policy": "中性",
            "overseas_impact": "有限",
            "data_freshness": "测试数据",
        },
        macro_score=60,
        macro_bias="中性",
    )
    risk = m.RiskScreeningReport(risk_score=80, risk_bias="可控")
    capital = m.CapitalFlowReport(
        northbound={"bias": "净流出"}, capital_score=50, capital_bias="中性"
    )
    review = m.ReviewResult(passed=True, score=90, summary="通过")

    outputs = [
        m.render_market_data_report(market),
        m.render_technical_report(technical),
        m.render_financial_report(financial),
        m.render_theme_report(theme),
        m.render_macro_report(macro),
        m.render_risk_report(risk),
        m.render_capital_report(capital),
        m.render_review_result(review),
    ]
    assert all(isinstance(output, str) and output.strip() for output in outputs)


def test_us_leveraged_signal_fails_closed_and_bounds_exposure():
    m = load_schemas()

    default_signal = m.LeveragedETFSignal()
    assert default_signal.direction == "avoid"
    assert default_signal.recommended == []

    with pytest.raises(ValueError, match="avoid direction"):
        m.LeveragedETFSignal(
            direction="avoid",
            recommended=[
                {
                    "ticker": "TQQQ",
                    "leverage": 3,
                    "position_pct": 5,
                    "stop_loss": -5,
                    "max_hold_days": 5,
                }
            ],
        )

    with pytest.raises(ValueError, match="at least one supported product"):
        m.LeveragedETFSignal(direction="long")

    with pytest.raises(ValueError, match="complete confirmed inputs"):
        m.LeveragedETFSignal(
            direction="long",
            recommended=[
                {
                    "ticker": "TQQQ",
                    "leverage": 3,
                    "position_pct": 3,
                    "stop_loss": -5,
                    "max_hold_days": 2,
                }
            ],
        )

    with pytest.raises(ValueError, match="unsupported leveraged ETF"):
        m.LeveragedETFPosition(
            ticker="FNGU",
            leverage=3,
            position_pct=2,
            stop_loss=-5,
            max_hold_days=2,
        )

    with pytest.raises(ValueError, match="match signal direction"):
        m.LeveragedETFSignal(
            direction="long",
            inputs_complete=True,
            trend_confirmed=True,
            momentum_confirmed=True,
            liquidity_confirmed=True,
            volatility_confirmed=True,
            vix_value=20,
            recommended=[
                {
                    "ticker": "SQQQ",
                    "leverage": 3,
                    "position_pct": 3,
                    "stop_loss": -5,
                    "max_hold_days": 2,
                }
            ],
        )

    with pytest.raises(ValueError, match="unsupported leveraged ETF"):
        m.LeveragedETFSignal(
            direction="long",
            inputs_complete=True,
            trend_confirmed=True,
            momentum_confirmed=True,
            liquidity_confirmed=True,
            volatility_confirmed=True,
            vix_value=20,
            recommended=[
                {
                    "ticker": "XYZ",
                    "leverage": 3,
                    "position_pct": 3,
                    "stop_loss": -5,
                    "max_hold_days": 2,
                }
            ],
        )

    with pytest.raises(ValueError, match="target leverage"):
        m.LeveragedETFSignal(
            direction="long",
            inputs_complete=True,
            trend_confirmed=True,
            momentum_confirmed=True,
            liquidity_confirmed=True,
            volatility_confirmed=True,
            vix_value=20,
            recommended=[
                {
                    "ticker": "TQQQ",
                    "leverage": 0.1,
                    "position_pct": 3,
                    "stop_loss": -5,
                    "max_hold_days": 2,
                }
            ],
        )

    with pytest.raises(ValueError, match="max_hold_days"):
        m.LeveragedETFPosition(
            ticker="SQQQ",
            leverage=3,
            position_pct=3,
            stop_loss=-5,
            max_hold_days=5,
        )

    with pytest.raises(ValueError, match="stop_loss must equal"):
        m.LeveragedETFPosition(
            ticker="TQQQ",
            leverage=3,
            position_pct=3,
            stop_loss=-3,
            max_hold_days=5,
        )

    with pytest.raises(ValueError, match="position_pct"):
        m.LeveragedETFPosition(
            ticker="TQQQ",
            leverage=3,
            position_pct=0,
            stop_loss=-5,
            max_hold_days=5,
        )

    with pytest.raises(ValueError, match="notional exposure"):
        m.LeveragedETFSignal(
            direction="long",
            inputs_complete=True,
            trend_confirmed=True,
            momentum_confirmed=True,
            liquidity_confirmed=True,
            volatility_confirmed=True,
            vix_value=20,
            recommended=[
                {
                    "ticker": "TQQQ",
                    "leverage": 3,
                    "position_pct": 12,
                    "stop_loss": -5,
                    "max_hold_days": 5,
                },
                {
                    "ticker": "SOXL",
                    "leverage": 3,
                    "position_pct": 8,
                    "stop_loss": -5,
                    "max_hold_days": 4,
                },
            ],
        )

    with pytest.raises(ValueError, match="VIX above 30"):
        m.LeveragedETFSignal(
            direction="inverse",
            inputs_complete=True,
            trend_confirmed=True,
            momentum_confirmed=True,
            liquidity_confirmed=True,
            volatility_confirmed=True,
            vix_value=31,
            recommended=[
                {
                    "ticker": "SQQQ",
                    "leverage": 3,
                    "position_pct": 3,
                    "stop_loss": -5,
                    "max_hold_days": 2,
                }
            ],
        )


def test_us_models_do_not_share_mutable_defaults():
    m = load_schemas()
    first = m.USFinancialReport(ticker="AAA")
    second = m.USFinancialReport(ticker="BBB")
    first.data_sources.append("SEC")
    first.peer_comparison.append({"ticker": "PEER"})
    assert second.data_sources == []
    assert second.peer_comparison == []

    risk_a = m.USRiskReport(ticker="AAA")
    risk_b = m.USRiskReport(ticker="BBB")
    risk_a.warnings.append("test")
    assert risk_b.warnings == []


def test_schema_renderers_accept_partial_optional_market_data():
    m = load_schemas()
    macro = m.MacroAnalysisReport(
        global_markets={"dji": {"close": 100.0}},
        macro_analysis={
            "monetary_policy": "中性",
            "fiscal_policy": "中性",
            "overseas_impact": "未知",
            "data_freshness": "部分缺失",
        },
        macro_score=50,
        macro_bias="中性",
    )
    assert "DJI: 100.0 (N/A)" in m.render_macro_report(macro)

    market = m.MarketDataReport(
        emotion={"score": 50, "level": "中性"},
        sector_flow={},
        market_filter={"level": "震荡", "position_pct": 40},
        northbound={"total_yi": 0, "bias": "持平"},
    )
    assert "北向资金**: 0.0亿" in m.render_market_data_report(market)


def test_health_check_uses_profile_aware_hermes_home():
    script = ROOT / "scripts/check_references_health.py"
    source = script.read_text()
    assert 'os.environ.get("HERMES_HOME"' in source
    env = os.environ.copy()
    env["HERMES_HOME"] = "/tmp/pantalone-profile-test"
    result = subprocess.run(
        [
            sys.executable,
            "-c",
            (
                "import importlib.util; "
                f"s=importlib.util.spec_from_file_location('health', {str(script)!r}); "
                "m=importlib.util.module_from_spec(s); s.loader.exec_module(m); "
                "print(m.INSTALLED_SCRIPTS_DIR)"
            ),
        ],
        check=True,
        capture_output=True,
        text=True,
        env=env,
    )
    assert result.stdout.strip() == "/tmp/pantalone-profile-test/scripts/amadeus"


def test_health_check_is_independent_of_checkout_depth(tmp_path):
    script = ROOT / "scripts/check_references_health.py"
    source = script.read_text()
    assert re.search(r"parents\[\d+\]", source) is None

    spec = importlib.util.spec_from_file_location("pantalone_health_depth", script)
    assert spec is not None and spec.loader is not None
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)

    shallow = tmp_path / "clone"
    shallow.mkdir()
    assert module.find_ancestor_scripts_dir(shallow) is None

    host = tmp_path / "host"
    expected = host / "scripts" / "amadeus"
    expected.mkdir(parents=True)
    deep = host / "a" / "b" / "clone"
    deep.mkdir(parents=True)
    assert module.find_ancestor_scripts_dir(deep) == expected


def test_structured_output_contract_does_not_claim_automatic_wiring():
    protocol = (ROOT / "subagents/protocol.md").read_text()
    readme = (ROOT / "subagents/README.md").read_text()
    combined = protocol + "\n" + readme
    assert "所有subagent输出采用" not in combined
    assert "通过 `with_structured_output(Schema)` 产生" not in combined
    assert "Schema是可选契约" in combined


def test_active_templates_are_consistent_and_profile_safe():
    template_paths = sorted((ROOT / "templates").rglob("*.md"))
    findings = []
    for path in template_paths:
        text = path.read_text()
        if "~/" in text or ("/" + "home" + "/") in text or ("/" + "root" + "/") in text:
            findings.append(str(path.relative_to(ROOT)))
        assert text.count("**信心度**: {高/中/低}") <= 1
        assert "**置信度**: {0.xx}（{H/M/L}）\n**置信度**" not in text
    assert findings == []

    strategy = (ROOT / "templates/strategy-template.md").read_text()
    assert "A+/A/B/C" in strategy


def test_tracked_template_backups_are_not_part_of_the_release_tree():
    backups = [
        str(path.relative_to(ROOT))
        for path in tracked_text_files()
        if "templates" in path.parts and ".bak" in path.name
    ]
    assert backups == []


def test_delivery_contract_keeps_word_optional():
    skill = (ROOT / "SKILL.md").read_text()
    review = (ROOT / "subagents/review.md").read_text()
    combined = skill + "\n" + review
    assert "按渠道长度和用户要求选择完整正文或Word" in skill
    assert "不得把 Word 作为所有报告的默认硬门槛" in review
    assert "完整版是否通过 .docx + MEDIA" not in combined


def test_historical_snapshots_are_excluded_from_the_release_tree():
    assert not (ROOT / "workflow.md").exists()
    assert not (ROOT / "SOUL.md.bak.20260514").exists()
    assert not (ROOT / "references/skill-full-reference.md").exists()
    backups = ROOT / "backups"
    assert not backups.exists() or not any(path.is_file() for path in backups.rglob("*"))
    assert not any(path.is_file() for path in (ROOT / "templates").rglob("*.bak*"))

    skill = (ROOT / "SKILL.md").read_text()
    workflow = (ROOT / "workflow_v4_unified.md").read_text()
    assert "历史全集、旧工作流和仓内备份不进入发布树" in skill
    assert "旧版本流程只保留在Git历史中" in workflow


def test_workflow_contract_remains_the_primary_entry():
    skill = (ROOT / "SKILL.md").read_text()
    assert "workflow_v4_unified.md" in skill
    assert "四层投资框架" in skill
    assert "8阶段" in skill
    assert "references/workflow-registry.yaml" in skill


def test_active_local_document_references_exist():
    paths = [
        *ACTIVE_DOCS,
        *CORE_REFERENCES,
        ROOT / "references/external-capabilities.yaml",
        ROOT / "references/workflow-registry.yaml",
    ]
    missing = []
    local_ref = re.compile(
        r"(?<!\$HERMES_HOME/)(?<![A-Za-z0-9_-])"
        r"((?:references|rules|templates|subagents)/[A-Za-z0-9_.-]+(?:/[A-Za-z0-9_.-]+)*\.(?:md|py|yaml))"
    )
    for source in paths:
        for target in local_ref.findall(source.read_text()):
            if not (ROOT / target).is_file():
                missing.append((str(source.relative_to(ROOT)), target))
    assert missing == []


def test_active_research_agent_does_not_advertise_removed_local_engine():
    research_agent = (ROOT / "subagents/research_agent.md").read_text()
    assert "amadeus_research.py" not in research_agent
    assert "research_templates/" not in research_agent
    assert "由Hermes会话按需编排" in research_agent


def test_readme_only_advertises_shipped_local_scripts():
    readme = (ROOT / "README.md").read_text()
    referenced = sorted(set(re.findall(r"(?<!HERMES_HOME/)scripts/([A-Za-z0-9_.-]+\.py)", readme)))
    missing = [name for name in referenced if not (ROOT / "scripts" / name).is_file()]
    assert missing == []


def test_readme_does_not_claim_unshipped_production_capabilities():
    readme = (ROOT / "README.md").read_text()
    forbidden = [
        "AUC 0.6512",
        "727 stocks",
        "606K samples",
        "v5.1 (prod)",
        "Current best",
        "Automated observation pool management",
        "OCIFQ (70%) + ML signal (30%)",
    ]
    assert [item for item in forbidden if item in readme] == []
    assert "A+/A/B/C" in readme
    assert "optional host capabilities" in readme
    assert "explicit authorization" in readme


def test_pool_rules_match_the_parent_a_plus_contract():
    rules = (ROOT / "rules/pool_rules.md").read_text()
    assert "A+" in rules
    assert "A+池 -10%" in rules
    assert "A池 -10%" in rules
    assert "B池 -5%" in rules
    assert "C池 -3%" in rules
    assert "A+池 180天" in rules
    assert "A池 180天" in rules
    assert "B池 60天" in rules
    assert "C池 30天" in rules
    assert "A池-8%" not in rules
    assert "C池>15天" not in rules


def test_stock_risk_rules_defer_to_authoritative_pool_stops_and_read_only_boundary():
    risk = (ROOT / "rules/risk_rules.md").read_text()
    assert "rules/pool_rules.md" in risk
    assert "A+/A/B/C" in risk
    fixed_stock_stops = [
        line
        for line in risk.splitlines()
        if "止损" in line and re.search(r"-\d+(?:\.\d+)?%", line)
    ]
    assert fixed_stock_stops == []
    assert "默认只生成待执行建议" in risk
    assert "用户明确授权" in risk
    assert "任一条件触发即执行" not in risk
    assert "无条件清仓" not in risk

    unauthorized_fixed_stops = []
    authoritative = {ROOT / "SKILL.md", ROOT / "rules/pool_rules.md"}
    for path in ACTIVE_DOCS:
        for line_number, line in enumerate(path.read_text().splitlines(), 1):
            if "止损" not in line or not re.search(r"-\d+(?:\.\d+)?%", line):
                continue
            if path in authoritative or "ETF" in line:
                continue
            unauthorized_fixed_stops.append((str(path.relative_to(ROOT)), line_number, line))
    assert unauthorized_fixed_stops == []

    current_rules = risk.split("### 历史验证", 1)[0]
    take_profit_levels = set(re.findall(r"\+(\d+)%", current_rules))
    assert {"12", "20", "30"}.issubset(take_profit_levels)
    assert not ({"8", "15", "25"} & take_profit_levels)


def test_active_execution_surface_requires_advisory_language_and_explicit_authorization():
    forbidden = re.compile(
        r"无条件退池|立即买入(?:/加仓)?|立即加仓|积极买入|重仓持有|直接退池|"
        r"必须执行\s*[:：]\s*\{?(?:止损|止盈|减仓)|执行变更"
    )
    findings = []
    for path in ACTIVE_DOCS:
        for line_number, line in enumerate(path.read_text().splitlines(), 1):
            if forbidden.search(line):
                findings.append((str(path.relative_to(ROOT)), line_number, line))
    assert findings == []


def test_active_templates_have_no_hard_line_count_gate():
    hard_line_gate = re.compile(
        r"(?:\d+\s*[-–]\s*\d+|不超过\s*\d+|至少\s*\d+|不少于\s*\d+)\s*行"
    )
    findings = []
    for path in sorted((ROOT / "templates").glob("*.md")):
        for line_number, line in enumerate(path.read_text().splitlines(), 1):
            if hard_line_gate.search(line):
                findings.append((str(path.relative_to(ROOT)), line_number, line))
    assert findings == []

    write_commands = re.compile(r"pool_manager\.py\s+(?:auto|apply|add|remove)\b")
    unsafe_calls = []
    for path in sorted((ROOT / "templates").glob("*.md")):
        for line_number, line in enumerate(path.read_text().splitlines(), 1):
            if write_commands.search(line) and "禁止" not in line:
                unsafe_calls.append((str(path.relative_to(ROOT)), line_number, line))
    assert unsafe_calls == []


def test_retained_core_references_are_portable_and_truthful():
    private_patterns = [
        re.escape("~/" + ".hermes"),
        re.escape("/" + "home" + "/" + "ubuntu"),
        re.escape("/" + "root" + "/"),
        re.escape("/" + "Users" + "/"),
    ]
    findings = []
    insecure_urls = []
    for path in CORE_REFERENCES:
        assert path.is_file(), path
        text = path.read_text()
        for line_number, line in enumerate(text.splitlines(), 1):
            for pattern in private_patterns:
                if re.search(pattern, line):
                    findings.append((str(path.relative_to(ROOT)), line_number, pattern))
            plaintext_scheme = "http" + "://"
            for url in re.findall(re.escape(plaintext_scheme) + r"[^\s\"'`)]+", line):
                if not re.match(
                    re.escape(plaintext_scheme)
                    + r"(?:127\.0\.0\.1|localhost|\[::1\])(?::\d+)?(?:/|$)",
                    url,
                ):
                    insecure_urls.append((str(path.relative_to(ROOT)), line_number, url))
    assert findings == []
    assert insecure_urls == []

    write_docs = [
        path for path in CORE_REFERENCES
        if re.search(r"--(?:apply|write|buy|sell)\b|\bauto\b", path.read_text())
    ]
    for path in write_docs:
        assert "明确授权" in path.read_text(), path

    unsafe_deserialization = []
    for path in CORE_REFERENCES:
        for line_number, line in enumerate(path.read_text().splitlines(), 1):
            if re.search(r"\bpickle\.(?:load|loads)\s*\(", line):
                unsafe_deserialization.append((str(path.relative_to(ROOT)), line_number, line))
    assert unsafe_deserialization == []


def test_schema_defaults_and_review_invariants():
    m = load_schemas()
    macro = m.MacroAnalysisReport(macro_score=50, macro_bias="中性")
    capital = m.CapitalFlowReport(capital_score=50, capital_bias="中性")
    assert "数据缺失" in m.render_macro_report(macro)
    assert "N/A" in m.render_capital_report(capital)

    critical = {
        "type": "risk",
        "severity": "critical",
        "description": "阻断",
        "fix": "修复",
    }
    major = {**critical, "severity": "major"}
    with pytest.raises(ValueError):
        m.ReviewResult(passed=True, issues=[critical], score=90, summary="错误通过")
    with pytest.raises(ValueError):
        m.ReviewResult(passed=True, issues=[], score=69, summary="低分通过")
    with pytest.raises(ValueError):
        m.ReviewResult(passed=True, issues=[major], score=90, summary="重大问题错误通过")
    with pytest.raises(ValueError):
        m.ETFReviewResult(
            passed=True,
            etf_type="宽基",
            issues=[critical],
            conclusion="通过",
        )
    with pytest.raises(ValueError):
        m.ETFReviewResult(
            passed=True,
            etf_type="宽基",
            conclusion="退回修改",
        )
    with pytest.raises(ValueError):
        m.ETFReviewResult(
            passed=True,
            etf_type="宽基",
            issues=[major],
            conclusion="通过",
        )


def test_missing_data_defaults_fail_closed():
    m = load_schemas()
    assert m.MarketDataReport.model_fields["data_quality"].default == m.DataQuality.D
    stock = m.StockTechnical()
    technical = m.TechnicalAnalysisReport(stocks={})
    index = m.MarketIndex()
    screening = m.ScreeningResult()
    verdict = m.RiskVerdict()
    assert stock.source == ""
    assert technical.data_sources == []
    assert index.close is None and index.pct is None and index.source == ""
    assert all(
        value == m.RiskFlag.UNKNOWN
        for value in [
            screening.st_risk,
            screening.pledge_risk,
            screening.goodwill_risk,
            screening.audit_risk,
            screening.cashflow_risk,
            verdict.overall,
        ]
    )

    provenance_fields = []
    models = [
        value
        for value in vars(m).values()
        if isinstance(value, type) and issubclass(value, m.BaseModel) and value is not m.BaseModel
    ]
    for model in models:
        for field_name, field in model.model_fields.items():
            if field_name not in {"source", "data_sources", "provenance"}:
                continue
            default = field.default_factory() if field.default_factory is not None else field.default
            provenance_fields.append((model.__name__, field_name, default))
    assert provenance_fields
    assert all(default in (None, "") or default == [] for _, _, default in provenance_fields)


def test_all_schema_definitions_and_default_factories_are_valid():
    m = load_schemas()
    models = []
    for value in vars(m).values():
        if isinstance(value, type) and issubclass(value, m.BaseModel) and value is not m.BaseModel:
            models.append(value)
    assert models
    for model in models:
        schema = model.model_json_schema()
        assert schema.get("title")
        for field_name, field in model.model_fields.items():
            if field.default_factory is not None:
                try:
                    field.default_factory()
                except Exception as exc:
                    pytest.fail(f"{model.__name__}.{field_name} default_factory failed: {exc}")


def test_router_uses_runtime_capabilities_not_stale_model_state():
    router = (ROOT / "router.md").read_text()
    stale = ["mimo", "小米模型", "DeepSeek", "deepseek-v4-pro", "config已设"]
    assert [term for term in stale if term in router] == []
    assert re.search(r"余额\s*[：:=]?\s*\d", router) is None
    assert "运行时配置" in router
    assert "实际模型" in router


def test_specialized_sop_cannot_override_safety_or_reintroduce_size_gates():
    workflow = (ROOT / "workflow_v4_unified.md").read_text()
    sop = (ROOT / "references/deep-stock-research-unified.md").read_text()
    assert "安全、证据、合规和写入边界" in workflow
    assert workflow.index("安全、证据、合规和写入边界") < workflow.index("任务专用SOP")
    assert re.search(r"\d+\s*[-–]\s*\d+\s*行|\d+行", sop) is None
    assert "仅在用户明确授权" in sop
    assert "默认只生成待写入建议" in sop
    assert "optional" in sop and "capability probe" in sop and "fallback" in sop

    hard_line_gate = re.compile(r"(?:\d+\s*[-–]\s*\d+|至少\s*\d+|不少于\s*\d+)\s*行")
    findings = [str(path.relative_to(ROOT)) for path in CORE_REFERENCES if hard_line_gate.search(path.read_text())]
    assert findings == []


def test_a_plus_rating_and_pool_mapping_are_consistent():
    skill = (ROOT / "SKILL.md").read_text()
    pool = (ROOT / "rules/pool_rules.md").read_text()
    sector = (ROOT / "references/sector-screening-workflow.md").read_text()
    assert "A+仅由OCIFQ评级产生" in skill
    assert "≥85分：A+评级候选" in sector
    assert "不得自动入池" in sector
    assert "建议入B池" not in sector


def test_optional_capability_probe_fails_closed_without_host_install(tmp_path):
    host_home = tmp_path / "empty-hermes-home"
    host_home.mkdir()
    result = subprocess.run(
        [
            sys.executable,
            str(ROOT / "scripts/probe_external_capabilities.py"),
            "--hermes-home",
            str(host_home),
        ],
        check=True,
        capture_output=True,
        text=True,
    )
    payload = __import__("json").loads(result.stdout)
    assert payload["total"] > 0
    assert payload["available"] == 0
    assert payload["missing_optional"] == payload["total"]
    assert payload["required_missing"] == []
    assert all(item["fallback"] for item in payload["capabilities"])
    assert list(host_home.iterdir()) == []


def test_release_metadata_matches_skill_contract_version():
    project = tomllib.loads((ROOT / "pyproject.toml").read_text())
    skill = (ROOT / "SKILL.md").read_text()
    skill_version = re.search(r"^version:\s*[\"']?([^\s\"']+)", skill, re.MULTILINE)
    assert skill_version is not None
    assert project["tool"]["pantalone"]["version"] == skill_version.group(1)
    assert project["tool"]["pantalone"]["distribution"] is False
    assert "build-system" not in project
    assert "project" not in project


def test_health_check_reports_static_scope_and_exclusions():
    result = subprocess.run(
        [sys.executable, str(ROOT / "scripts/check_references_health.py"), "--quiet"],
        check=True,
        capture_output=True,
        text=True,
    )
    assert "所列静态检查通过" in result.stdout
    assert "未覆盖：" in result.stdout


def test_all_tracked_files_reject_high_entropy_credential_literals():
    result = subprocess.run(
        ["git", "ls-files", "-z"],
        cwd=ROOT,
        check=True,
        capture_output=True,
    )
    credential = re.compile(
        rb"(?i)(?:^|[^A-Za-z0-9_])(?:token|api[_-]?key|secret)\b\s*[:=]\s*[`\"']?([0-9a-f]{40,})"
    )
    findings = []
    for raw_path in result.stdout.split(b"\0"):
        if not raw_path:
            continue
        path = ROOT / os.fsdecode(raw_path)
        if not path.is_file():
            continue
        data = path.read_bytes()
        if credential.search(data):
            findings.append(str(path.relative_to(ROOT)))
    assert findings == []


def test_all_tracked_text_is_portable_and_has_no_external_plaintext_http():
    private_markers = [
        "~/" + ".hermes",
        "/" + "home" + "/" + "ubuntu",
        "/" + "root" + "/",
        "/" + "Users" + "/",
    ]
    plaintext_scheme = "http" + "://"
    loopback = re.compile(
        r"^(?:127\.0\.0\.1|localhost|\[::1\])(?::\d+)?(?:/|$)"
    )
    hard_line_gate = re.compile(
        r"(?:\d+\s*[-–]\s*\d+|不超过\s*\d+|至少\s*\d+|不少于\s*\d+)\s*行"
    )
    wildcard_listener = "0" + ".0.0.0"
    findings = []
    for path in tracked_text_files():
        relative = str(path.relative_to(ROOT))
        for line_number, line in enumerate(path.read_text(encoding="utf-8").splitlines(), 1):
            for marker in private_markers:
                if marker in line:
                    findings.append((relative, line_number, "private_path"))
            for url in re.findall(re.escape(plaintext_scheme) + r"[^\s\"'`)]+", line):
                host_and_path = url[len(plaintext_scheme):]
                if not loopback.match(host_and_path):
                    findings.append((relative, line_number, "external_plaintext_http"))
            if hard_line_gate.search(line):
                findings.append((relative, line_number, "hard_line_gate"))
            if wildcard_listener in line:
                findings.append((relative, line_number, "wildcard_listener"))
    assert findings == []


def test_all_tracked_docs_use_authoritative_stock_risk_contract():
    old_a_stop = "A池" + "-8%"
    old_c_timeout = "C池" + ">15天"
    old_b_timeout = "B池" + ">30天"
    high_price_override = re.compile(
        r"max\s*\([^\n,]*(?:0\.25|25%|pct_limit)[^\n,]*,\s*min_lot_cost\s*\)|"
        r"确保高价股能买至少1手|max\s*\(\s*25%\s*,\s*1手\s*\)|"
        r"max\s*\([^\n]*25%[^\n]*当前价格[×*]100[^\n]*\)"
    )
    unconditional_execution = re.compile(
        r"任一条件触发即执行|无条件(?:减半仓|清仓|退池)|直接退池"
    )
    alternative_exit_sop = re.compile(
        r"(?:跌破|下穿)\s*(?:MA|均线)?\s*(?:5|10)[^\n]{0,25}(?:减仓|清仓|卖出)|"
        r"(?:5\s*[-–]\s*7天|10天\+?|持仓超过(?:5|10)天)[^\n]{0,25}(?:减仓|清仓|卖出)|"
        r"三层(?:清仓|退出)"
    )
    moving_average_evidence = re.compile(
        r"(?<![A-Za-z0-9_])MA(?:5|10|20|50|60|120)?(?![A-Za-z0-9_])|均线",
        re.IGNORECASE,
    )
    stock_action = re.compile(r"买入|卖出|减仓|清仓|止损|止盈|退出")
    evidence_only = re.compile(r"不得|不能|禁止|仅作|不生成|不应|不可|非.{0,8}触发")
    technical_heading = re.compile(
        r"^#{2,4}\s.*(?:\bMA(?:5|10|20|50|60|120)?\b|均线|RSI|MACD|BOLL|布林带|量价|技术(?:指标|形态|信号|面))",
        re.IGNORECASE,
    )
    action_field = re.compile(
        r"(?:\*\*(?:信号|结论)\*\*|(?:信号|结论)\s*[:|])[^\n]{0,160}(?:买入|卖出|减仓|清仓|止损|止盈|退出)"
    )
    rating_exit = re.compile(r"\|\s*\*\*[DF]\*\*\s*\|[^\n]*(?:减仓|清仓)")
    fixed_exit_sop = re.compile(
        r"(?:Sell:\s*Score|Stop-loss:\s*-\d|Take-profit:\s*\+\d|Max holding:\s*\d)|"
        r"(?:二级止损|三级强平)[^\n]{0,80}(?:减仓|清仓|卖出|退出)",
        re.IGNORECASE,
    )
    score_direct_action = re.compile(
        r"(?:ML评分|ML_SCORE|ML score|composite_score)[^\n]{0,100}(?:买入|卖出|减仓|清仓|加仓|强平)|"
        r"(?:买入|卖出|减仓|清仓|加仓|强平)[^\n]{0,100}(?:ML评分|ML_SCORE|ML score|composite_score)",
        re.IGNORECASE,
    )
    single_signal_action = re.compile(
        r"(?:温度骤[升降]|版本.{0,8}(?:确认|高潮|退潮)|情绪.{0,5}(?:回暖|降温|冰点))[^\n]{0,100}(?:加仓|减仓|清仓|卖出|止损|止盈)|"
        r"(?:加仓|减仓|清仓|卖出|止损|止盈)[^\n]{0,100}(?:温度骤[升降]|版本.{0,8}(?:确认|高潮|退潮)|情绪.{0,5}(?:回暖|降温|冰点))|"
        r"\b(?:buy_threshold|sell_threshold)\b[^\n]{0,20}(?:买入|卖出)",
        re.IGNORECASE,
    )
    findings = []
    for path in tracked_text_files():
        if path.suffix.lower() != ".md":
            continue
        text = path.read_text(encoding="utf-8")
        relative = str(path.relative_to(ROOT))
        if old_a_stop in text:
            findings.append((relative, old_a_stop))
        if old_c_timeout in text:
            findings.append((relative, old_c_timeout))
        if old_b_timeout in text:
            findings.append((relative, old_b_timeout))
        if high_price_override.search(text):
            findings.append((relative, "high_price_override"))
        if unconditional_execution.search(text):
            findings.append((relative, "unconditional_execution"))
        if alternative_exit_sop.search(text):
            findings.append((relative, "alternative_exit_sop"))
        if rating_exit.search(text):
            findings.append((relative, "rating_driven_exit"))
        if fixed_exit_sop.search(text):
            findings.append((relative, "fixed_exit_sop"))
        if score_direct_action.search(text):
            findings.append((relative, "score_direct_action"))
        if single_signal_action.search(text):
            findings.append((relative, "single_signal_action"))
        for line_number, line in enumerate(text.splitlines(), 1):
            if (
                moving_average_evidence.search(line)
                and stock_action.search(line)
                and not evidence_only.search(line)
            ):
                findings.append((relative, line_number, "moving_average_action_mapping"))
        lines = text.splitlines()
        heading_indexes = [i for i, line in enumerate(lines) if re.match(r"^#{2,4}\s", line)]
        for index, start in enumerate(heading_indexes):
            end = heading_indexes[index + 1] if index + 1 < len(heading_indexes) else len(lines)
            if technical_heading.search(lines[start]) and action_field.search("\n".join(lines[start:end])):
                findings.append((relative, start + 1, "technical_block_action_field"))
    assert findings == []


def test_building_strategy_defers_to_authoritative_exit_rules():
    text = (ROOT / "references/building-strategy-framework.md").read_text()
    assert "rules/pool_rules.md" in text
    assert "rules/risk_rules.md" in text
    assert "+12%" in text and "+20%" in text and "+30%" in text
    conflicting = ["+10-15%", "+20-25%", "+40%+", "跌破MA60", "跌破MA20"]
    assert [value for value in conflicting if value in text] == []
    direct_entry_mappings = [
        "突破确认后买入",
        "消息确认后快速建仓",
        "根据走势决定加仓",
        "完成底仓建设",
        "底仓 | 30-40%",
    ]
    assert [value for value in direct_entry_mappings if value in text] == []
    assert "完整建仓门槛" in text
    assert "待授权建仓建议" in text


def test_active_market_and_etf_evidence_cannot_directly_change_positions():
    trading = (ROOT / "rules/trading_rules.md").read_text(encoding="utf-8")
    morning = (ROOT / "templates/morning_template.md").read_text(encoding="utf-8")
    etf = (ROOT / "subagents/etf.md").read_text(encoding="utf-8")

    assert "高波动，缩小仓位或不触发" not in trading
    assert "不得据此单独缩仓" in trading
    assert "降低仓位，不追涨" not in morning
    assert "可小仓位试探" not in morning
    assert "不得单独触发降仓" in morning
    assert "不得单独触发试探建仓" in morning
    assert "冲高5%-8%开始减仓" not in etf
    assert "任何模拟盘写入必须获得用户对该动作的逐次明确授权" in etf
    assert "减仓、清仓或卖出必须获得用户对该动作的逐次明确授权" in etf


def test_scores_positions_and_leader_labels_cannot_bypass_authorization():
    score_to_action = re.compile(
        r"(?:OCIFQ|技术面?|tech|总分|综合分|任意评分|评分|得分|score)[^\n]{0,120}"
        r"(?:立即建仓|建议建仓|可建仓|买入|卖出|仓位减半|入\s*\d+\s*%)",
        re.IGNORECASE,
    )
    oversized_position = re.compile(
        r"底仓\s*(?:\||[:：])?\s*30\s*[-–—~～至]\s*40\s*%|"
        r"(?:底仓|波段仓|机动仓)[^\n]{0,30}(?:30\s*[-–—~～至]\s*40|20\s*[-–—~～至]\s*30)\s*%"
    )
    automatic_pool_write = re.compile(
        r"(?:自动|直接|默认)[^\n]{0,12}(?:入观察池|入池)|"
        r"(?:版本龙头|板块龙头|各板块龙头)[^\n]{0,20}(?:自动|直接|默认)?入观察池"
    )
    phase_to_position = re.compile(
        r"(?:潜伏期|确认期|加速期|高潮期|退潮期)[^\n]{0,100}"
        r"(?:建立底仓|加仓|减仓|清仓|改变仓位|分批止盈)"
    )
    allowed = re.compile(r"不得|不能|禁止|仅作|不生成|待授权|明确授权|证据|不直接")
    findings = []
    for path in tracked_text_files():
        if path.suffix.lower() != ".md" or "tests" in path.parts:
            continue
        relative = str(path.relative_to(ROOT))
        for line_number, line in enumerate(path.read_text(encoding="utf-8").splitlines(), 1):
            if allowed.search(line):
                continue
            for name, pattern in (
                ("score_to_action", score_to_action),
                ("oversized_position", oversized_position),
                ("automatic_pool_write", automatic_pool_write),
                ("phase_to_position", phase_to_position),
            ):
                if pattern.search(line):
                    findings.append((relative, line_number, name))
    assert findings == []


def test_authoritative_entry_contract_has_no_score_or_market_position_mapping():
    risk = (ROOT / "rules/risk_rules.md").read_text(encoding="utf-8")
    checklist = (ROOT / "subagents/checklist.md").read_text(encoding="utf-8")
    forbidden = [
        "总分≥60分才可生成建仓建议",
        "60-69分 | 10%",
        "70-79分 | 15%",
        "80-89分 | 20%",
        "90+分 | 25%",
        "首批50%计划仓位建议",
        "剩余50%计划仓位建议",
        "可生成积极建仓建议",
        "可生成常规建仓建议",
        "评分最高的优先",
        "买入计划是否通过评分系统（≥60分）",
        "min(25%总仓位, 5万)",
    ]
    joined = risk + "\n" + checklist
    assert [value for value in forbidden if value in joined] == []
    assert "完整建仓资格门槛" in risk
    assert "不得按评分档位套用10%/15%/20%/25%的固定仓位" in risk
    assert "不使用固定50/50" in risk
    assert "同一股票的底仓、波段仓、机动仓和其他名义仓位合计不得超过25%" in risk
    assert "评分是否仅作研究证据" in checklist
    assert "8项建仓资格门槛" in checklist


def test_documented_state_writes_require_per_action_authorization():
    write_command = re.compile(
        r"(?:pool_manager|amadeus_pool_manager|amadeus_etf_pool_manager|"
        r"amadeus_sim_integrate)\.py\s+"
        r"(?:auto|apply|add|remove|record|daily_update)\b"
    )
    findings = []
    for path in tracked_text_files():
        text = path.read_text(encoding="utf-8")
        if write_command.search(text) and "逐次明确授权" not in text:
            findings.append(str(path.relative_to(ROOT)))
    assert findings == []


def test_yolo_memory_and_cron_authorization_boundaries_are_explicit():
    soul = (ROOT / "SOUL.md").read_text(encoding="utf-8")
    yolo_start = soul.index("### 15.1")
    yolo_end = soul.index("### 15.2")
    yolo_section = soul[yolo_start:yolo_end]

    assert "记忆写入" not in yolo_section
    assert "Cron Job创建时" in soul
    assert "限域预授权" in soul
    assert "范围内定时触发无需重复确认" in soul
    assert "任何范围变更必须重新授权" in soul
    assert "自动推送视为预授权，不需逐次确认" not in soul


def test_simulator_uses_profile_home_and_authoritative_pool_stops(monkeypatch, tmp_path):
    monkeypatch.setenv("HERMES_HOME", str(tmp_path))
    module = load_sim_integrator()
    assert module.DB_PATH == tmp_path / "cache/amadeus/simulator.db"
    assert module.calculate_stop_loss(100.0, "A+") == 90.0
    assert module.calculate_stop_loss(100.0, "A") == 90.0
    assert module.calculate_stop_loss(100.0, "B") == 95.0
    assert module.calculate_stop_loss(100.0, "C") == 97.0
    with pytest.raises(ValueError):
        module.calculate_stop_loss(100.0, "UNKNOWN")


def test_simulator_write_commands_require_explicit_authorization():
    module = load_sim_integrator()
    assert module.write_authorized(["daily_update"]) is False
    assert module.write_authorized(["record", "{}"]) is False
    assert module.write_authorized(["daily_update", "--authorized"]) is True
    assert module.write_authorized(["record", "{}", "--authorized"]) is True


def test_simulator_write_functions_fail_closed_without_authorization(monkeypatch, tmp_path):
    monkeypatch.setenv("HERMES_HOME", str(tmp_path))
    module = load_sim_integrator()
    assert module.daily_update() == {"error": "explicit_authorization_required"}
    assert module.record_trade("{}") == {"error": "explicit_authorization_required"}
    assert not module.DB_PATH.exists()


def test_simulator_empty_profile_is_read_only_and_writes_validated(monkeypatch, tmp_path):
    monkeypatch.setenv("HERMES_HOME", str(tmp_path))
    module = load_sim_integrator()

    status = module.get_status()
    assert status["total_value"] == module.INIT_CAPITAL
    assert status["positions"] == []
    assert not module.DB_PATH.exists()

    invalid = [
        '{"action":"hold","code":"600000","price":10,"shares":100}',
        '{"action":"buy","code":"","price":10,"shares":100}',
        '{"action":"buy","code":"600000","price":0,"shares":100}',
        '{"action":"buy","code":"600000","price":10,"shares":1.5}',
        '{"action":"sell","price":10,"position_id":0}',
    ]
    for payload in invalid:
        result = module.record_trade(payload, authorized=True)
        assert "error" in result
        assert not module.DB_PATH.exists()

    created = module.record_trade(
        '{"action":"buy","code":"600000","name":"浦发银行","pool":"B","price":10,"shares":100}',
        authorized=True,
    )
    assert created["status"] == "recorded"
    assert module.DB_PATH.is_file()
    assert module.get_status()["positions_count"] == 1


def test_simulator_status_does_not_initialize_existing_empty_database(monkeypatch, tmp_path):
    monkeypatch.setenv("HERMES_HOME", str(tmp_path))
    module = load_sim_integrator()
    assert module.DB_PATH is not None
    module.DB_PATH.parent.mkdir(parents=True)
    module.DB_PATH.touch()

    before = module.DB_PATH.stat()
    assert module.get_status() == {"error": "simulator_schema_missing"}
    after = module.DB_PATH.stat()
    assert after.st_size == before.st_size == 0
    assert after.st_mtime_ns == before.st_mtime_ns


def test_simulator_rejects_explicit_empty_hermes_home(monkeypatch):
    monkeypatch.setenv("HERMES_HOME", "")
    module = load_sim_integrator()
    assert module.DB_PATH is None
    assert module.get_status() == {"error": "HERMES_HOME_must_be_nonempty"}
    result = module.record_trade(
        '{"action":"buy","code":"600000","price":10,"shares":100}',
        authorized=True,
    )
    assert result == {"error": "HERMES_HOME_must_be_nonempty"}


def test_simulator_sell_response_reports_actual_position_shares(monkeypatch, tmp_path):
    monkeypatch.setenv("HERMES_HOME", str(tmp_path))
    module = load_sim_integrator()
    bought = module.record_trade(
        '{"action":"buy","code":"600000","pool":"B","price":10,"shares":100}',
        authorized=True,
    )
    assert bought["shares"] == 100
    position_id = module.get_status()["positions"][0]["id"]
    sold = module.record_trade(
        json.dumps({"action": "sell", "position_id": position_id, "price": 11}),
        authorized=True,
    )
    assert sold["code"] == "600000"
    assert sold["shares"] == 100


def test_health_check_detects_markdown_colon_credentials_without_echoing_values(tmp_path):
    script = ROOT / "scripts/check_references_health.py"
    spec = importlib.util.spec_from_file_location("pantalone_health_credentials", script)
    assert spec is not None and spec.loader is not None
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)

    token = "a" * 56
    (tmp_path / "reference.md").write_text(f"token: `{token}`\n")
    setattr(module, "PANTALONE_ROOT", tmp_path)
    setattr(module, "SCRIPTS_DIR", tmp_path / "missing-amadeus")

    result = module.check_credentials()

    assert result["clean"] is False
    assert result["total"] == 1
    assert result["findings"] == [
        {"file": "pantalone/reference.md", "line": 1}
    ]
    assert token not in repr(result)


def test_md2docx_warns_when_malformed_table_falls_back(tmp_path):
    source = tmp_path / "malformed.md"
    output = tmp_path / "malformed.docx"
    source.write_text("| A | B |\n|---|---|\n| 1 | 2 | 3 |\n")
    result = subprocess.run(
        [sys.executable, str(ROOT / "scripts/md2docx.py"), str(source), str(output)],
        check=True,
        capture_output=True,
        text=True,
    )
    assert "WARNING: malformed markdown table" in result.stderr
    assert output.is_file()


def test_tencent_https_quote_parser_handles_gbk_batch_fixture():
    parser = load_tencent_quote_parser()

    def record(variable, name, code, price, change_pct):
        fields = [""] * 48
        fields[1] = name
        fields[2] = code
        fields[3] = str(price)
        fields[32] = str(change_pct)
        return f'{variable}="{"~".join(fields)}";'

    fixture = (
        record("v_hk00700", "腾讯控股", "00700", 350.4, 1.25)
        + "\n"
        + record("v_hk09988", "阿里巴巴-W", "09988", 81.55, -0.75)
    ).encode("gbk")
    assert parser.parse_tencent_quotes(fixture) == [
        {"code": "00700", "name": "腾讯控股", "price": 350.4, "change_pct": 1.25, "field_count": 48},
        {"code": "09988", "name": "阿里巴巴-W", "price": 81.55, "change_pct": -0.75, "field_count": 48},
    ]
    with pytest.raises(ValueError, match="expected at least 33 fields"):
        parser.parse_tencent_quotes(b'v_hk00001="x~y";')


@pytest.mark.skipif(
    os.environ.get("PANTALONE_LIVE_HTTPS") != "1",
    reason="opt-in read-only live HTTPS smoke test",
)
def test_tencent_https_live_smoke_is_read_only():
    import urllib.request

    parser = load_tencent_quote_parser()
    request = urllib.request.Request(
        "https://qt.gtimg.cn/q=hk00700",
        headers={"User-Agent": "Pantalone-read-only-smoke/1.0"},
    )
    with urllib.request.urlopen(request, timeout=20) as response:
        assert response.status == 200
        quotes = parser.parse_tencent_quotes(response.read())
    assert quotes[0]["code"] == "00700"
    assert quotes[0]["price"] > 0


def test_md2docx_preserves_heading_levels_and_empty_table_cells(tmp_path):
    source = tmp_path / "sample.md"
    output = tmp_path / "sample.docx"
    source.write_text(
        "# 标题\n\n### 三级\n\n#### 四级\n\n"
        "| A | B | C |\n|---|---|---|\n| 1 | | 3 |\n"
    )
    result = subprocess.run(
        [sys.executable, str(ROOT / "scripts/md2docx.py"), str(source), str(output)],
        check=True,
        capture_output=True,
        text=True,
    )
    assert output.is_file() and output.stat().st_size > 0, result.stdout
    doc = Document(output)
    styles = {p.text: getattr(p.style, "name", "") for p in doc.paragraphs if p.text}
    assert styles["三级"] == "Heading 2"
    assert styles["四级"] == "Heading 3"
    assert [[cell.text for cell in row.cells] for row in doc.tables[0].rows] == [
        ["A", "B", "C"],
        ["1", "", "3"],
    ]
